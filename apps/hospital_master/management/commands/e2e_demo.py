"""e2e_demo — drives the whole application through its HTTP API, one role at a time, and writes an Excel list of everything it tried."""
import io
import json
import secrets
import string
import traceback
from datetime import date, timedelta
from unittest import mock

from django.core.management.base import BaseCommand
from django.test import Client, override_settings

ROLES = ["doctor", "nurse", "front_desk", "lab_tech", "pharmacist"]


def _password():
    pool = string.ascii_letters + string.digits
    return "Aa1!" + "".join(secrets.choice(pool) for _ in range(10))


class Run:
    def __init__(self, tag):
        self.tag = tag
        self.rows = []
        self.client = Client()
        self.secrets = {}

    def call(self, area, role, what, method, url, *, token=None, body=None, expect=(200, 201), files=None, query=None, check=None):
        headers = {"HTTP_AUTHORIZATION": f"Bearer {token}"} if token else {}
        try:
            fn = getattr(self.client, method)
            if files is not None:
                resp = fn(url, {**(body or {}), **files}, **headers)
            elif method == "get":
                resp = fn(url, query or body or {}, **headers)
            else:
                resp = fn(url, json.dumps(body or {}), content_type="application/json", **headers)
            try:
                data = resp.json()
            except Exception:
                data = {}
            ok = resp.status_code in expect
            note = "" if ok else (str(data.get("message") or data) + " " + str(data.get("errors") or ""))[:200]
            if ok and check:
                verdict = check(data.get("data", data) if isinstance(data, dict) else data)
                if verdict is not True:
                    ok, note = False, f"unexpected content: {verdict}"
        except Exception as exc:
            resp, data, ok, note = None, {}, False, f"{type(exc).__name__}: {exc}"[:160]
            traceback.print_exc()
        self.rows.append([area, role, what, f"{method.upper()} {url}", "/".join(map(str, expect)),
                          resp.status_code if resp is not None else "-", "PASS" if ok else "FAIL", note])
        return data.get("data", data) if isinstance(data, dict) else data

    def note(self, area, role, what, ok, detail=""):
        self.rows.append([area, role, what, "(direct check)", "-", "-", "PASS" if ok else "FAIL", detail])
        return ok


class Command(BaseCommand):
    help = "Run the whole application end to end through its API and write an Excel report."

    def add_arguments(self, p):
        p.add_argument("--excel", default="e2e_report.xlsx")
        p.add_argument("--patients", type=int, default=2)
        p.add_argument("--staff-per-role", type=int, default=1)
        p.add_argument("--tag", default="")

    def handle(self, *a, **o):
        import logging
        logging.disable(logging.CRITICAL)
        tag = o["tag"] or secrets.token_hex(3)
        r = Run(tag)
        with override_settings(CELERY_TASK_ALWAYS_EAGER=True, CELERY_TASK_EAGER_PROPAGATES=False, ALLOWED_HOSTS=["*"]):
            try:
                self.scenario(r, o)
            except Exception:
                traceback.print_exc()
                r.note("harness", "-", "scenario aborted", False, traceback.format_exc()[-300:])
        self.excel(r, o["excel"])
        bad = [x for x in r.rows if x[6] == "FAIL"]
        self.stdout.write(f"{len(r.rows)} checks, {len(bad)} failed -> {o['excel']}")
        for x in bad:
            self.stdout.write(f"  FAIL {x[0]} | {x[1]} | {x[2]} | {x[5]} | {x[7]}")

    def scenario(self, r, o):
        from django.contrib.auth import get_user_model
        tag = r.tag
        # ---- platform admin ------------------------------------------------------------------
        User = get_user_model()
        pw_platform = _password()
        User.objects.filter(username=f"e2e_{tag}").delete()
        User.objects.create_superuser(f"e2e_{tag}", f"e2e_{tag}@example.com", pw_platform)
        d = r.call("Login", "platform admin", "Platform admin logs in", "post", "/api/v1/auth/login/platform/",
                   body={"username": f"e2e_{tag}", "password": pw_platform})
        plat = d.get("access") or d.get("access_token") or d.get("token")
        r.note("Login", "platform admin", "Token issued", bool(plat), str(list(d)[:6]))
        r.call("Platform", "platform admin", "Own identity (/me)", "get", "/api/v1/auth/me/", token=plat)

        # ---- hospital -------------------------------------------------------------------------
        mob = lambda: "9" + "".join(secrets.choice("0123456789") for _ in range(9))
        code = f"e2e-{tag}"
        admin_mobile = mob()
        h = r.call("Platform", "platform admin", "Create a hospital (database provisioned)", "post", "/api/v1/platform/hospitals/",
                   token=plat, body={"name": f"E2E Hospital {tag}", "hospital_code": code, "city": "Bengaluru", "state": "Karnataka",
                                     "admin_mobile": admin_mobile, "admin_email": f"admin.{tag}@example.com", "tier": "enterprise"})
        cred = (h or {}).get("credentials") or {}
        r.note("Platform", "platform admin", "Hospital admin credentials returned", bool(cred.get("temp_password")), str(list(cred)))
        r.call("Platform", "platform admin", "Hospital list", "get", "/api/v1/platform/hospitals/", token=plat)
        r.call("Platform", "platform admin", "Hospital detail", "get", f"/api/v1/platform/hospitals/{h['id']}/", token=plat)

        # ---- hospital admin -------------------------------------------------------------------
        d = r.call("Login", "hospital admin", "Hospital admin logs in (mobile + temporary password)", "post", "/api/v1/auth/login/staff/",
                   body={"mobile": admin_mobile, "password": cred.get("temp_password", "")})
        adm = d.get("access")
        r.call("Login", "hospital admin", "Permissions of the logged-in admin", "get", "/api/v1/auth/me/permissions/", token=adm)
        client_put = r.client.put
        logins = [("platform admin", f"e2e_{tag}", pw_platform), ("hospital admin", admin_mobile, cred.get("temp_password", ""))]

        # ---- hospital set-up by the admin -----------------------------------------------------
        br = r.call("Hospital setup", "hospital admin", "List branches", "get", "/api/v1/org/branches/", token=adm)
        branches = br if isinstance(br, list) else (br or {}).get("results", [])
        branch_id = branches[0]["id"] if branches else None
        if branch_id is None:
            b = r.call("Hospital setup", "hospital admin", "Create a branch", "post", "/api/v1/org/branches/", token=adm,
                       body={"name": "Main", "city": "Bengaluru", "address": "1 Test Road"})
            branch_id = (b or {}).get("id")
        dep = r.call("Hospital setup", "hospital admin", "Create a department", "post", "/api/v1/org/departments/", token=adm,
                     body={"name": "Paediatrics", "branch": branch_id, "branch_id": branch_id})
        r.call("Hospital setup", "hospital admin", "Hospital settings", "get", "/api/v1/org/settings/", token=adm)
        r.call("Hospital setup", "hospital admin", "Staff list", "get", "/api/v1/org/staff/", token=adm)

        staff = {}
        for role in ROLES:
            for i in range(o["staff_per_role"]):
                m = mob()
                inv = r.call("Hospital setup", "hospital admin", f"Invite a {role}", "post", "/api/v1/org/staff/invite/", token=adm,
                             body={"first_name": role.title().replace("_", " "), "last_name": f"E2E{i + 1}", "phone": m,
                                   "email": f"{role}{i + 1}.{tag}@example.com", "role": role, "branch_id": branch_id,
                                   **({"department_id": (dep or {}).get("id")} if role == "doctor" and dep else {})})
                temp = (inv or {}).get("temp_password") or ""
                tokn = (inv or {}).get("invite_token") or ""
                if tokn:
                    pw = _password()
                    r.call("Login", role, f"{role} sets own password from the invite", "post", "/api/v1/auth/setup-password/",
                           body={"invite_token": tokn, "new_password": pw, "confirm_password": pw})
                else:
                    pw = temp
                lg = r.call("Login", role, f"{role} logs in", "post", "/api/v1/auth/login/staff/", body={"mobile": m, "password": pw})
                staff.setdefault(role, []).append({"mobile": m, "token": (lg or {}).get("access"), "id": (inv or {}).get("id") or (inv or {}).get("staff_id")})
                logins.append((role, m, pw))

        # ---- doctor profile, schedule, billing catalogue --------------------------------------
        doc = staff["doctor"][0]
        r.call("Hospital setup", "hospital admin", "Doctor profile (specialisation, fee)", "post", f"/api/v1/org/staff/{doc['id']}/doctor-profile/",
               token=adm, body={"registration_no": f"KMC{tag}", "specialisation": "Pediatrics", "qualification": "MD", "experience_years": 8,
                                "consultation_fee": "500", "followup_fee": "300"}, expect=(200, 201))
        days = [{"day_of_week": dw, "is_available": True, "start_time": "09:00", "end_time": "17:00"} for dw in range(7)]
        r.call("Hospital setup", "hospital admin", "Doctor weekly schedule", "put", f"/api/v1/org/staff/{doc['id']}/schedule/", token=adm,
               body={"slot_duration_minutes": 15, "days": days}, expect=(200, 201))
        cats = r.call("Billing", "hospital admin", "List service categories", "get", "/api/v1/billing/service-categories/", token=adm)
        cats = cats if isinstance(cats, list) else (cats or {}).get("results", [])
        cat = cats[0] if cats else None
        r.call("Billing", "hospital admin", "Create a billable service", "post", "/api/v1/billing/services/", token=adm,
               body={"name": "Paediatric consultation", "code": "CONS1", "unit_price": "500", "tax_rate": "0", "category": (cat or {}).get("id")})
        # ---- patient portal: sign-up with e-mail OTP, login ------------------------------------
        pats = []
        for i in range(o["patients"]):
            pats.append(self.portal_patient(r, i, mob, tag, logins))
        for pt in pats:
            self.patient_portal(r, pt, h["id"], doc, tag)
        self.documents(r, pats[0])
        self.hospital_visit(r, pats, staff, adm, branch_id, tag, mob)
        self.more_flows(r, pats, staff, adm, branch_id, h["id"], doc, (dep or {}).get("id"))
        self.access_control(r, staff, adm, plat, pats, logins, h["id"])
        r.call("Billing", "hospital admin", "Payment modes list", "get", "/api/v1/billing/payment-modes/", token=adm)

    def portal_patient(self, r, i, mob, tag, logins):
        email, m, pw = f"patient{i + 1}.{tag}@example.com", mob(), _password()
        with mock.patch("core.messaging.otp.generate_code", return_value="123456"):
            r.call("Patient sign-up", f"patient {i + 1}", "Ask for an e-mail OTP", "post", "/api/v1/auth/otp/request/",
                   body={"purpose": "registration_patient", "identifier": email})
            v = r.call("Patient sign-up", f"patient {i + 1}", "Verify the OTP", "post", "/api/v1/auth/otp/verify/",
                       body={"purpose": "registration_patient", "identifier": email, "code": "123456"})
        reg = r.call("Patient sign-up", f"patient {i + 1}", "Create the account", "post", "/api/v1/portal/register/",
                     body={"action_token": (v or {}).get("action_token", ""), "full_name": f"Asha Rao{i + 1} {tag}", "mobile": m,
                           "password": pw, "gender": "female", "date_of_birth": str(date.today() - timedelta(days=700 + 40 * i))})
        awpid = (reg or {}).get("awpid")
        lg = r.call("Login", f"patient {i + 1}", "Patient logs in with AWPID", "post", "/api/v1/auth/login/patient/", body={"mobile": awpid, "password": pw})
        logins.append((f"patient {i + 1}", awpid, pw))
        return {"awpid": awpid, "mobile": m, "token": (lg or {}).get("access"), "pw": pw}

    _booked = []

    def patient_portal(self, r, pt, hospital_id, doc, tag):
        n, tk = pt["awpid"], pt["token"]
        who = f"patient {n[-4:]}"
        r.call("Patient portal", who, "Own profile", "get", "/api/v1/portal/profile/", token=tk)
        r.call("Patient portal", who, "Hospitals list", "get", "/api/v1/portal/hospitals/", token=tk)
        r.call("Patient portal", who, "Doctors at the hospital", "get", f"/api/v1/portal/hospitals/{hospital_id}/doctors/", token=tk)
        d = date.today() + timedelta(days=1)
        r.call("Patient portal", who, "Free slots for a doctor", "get", f"/api/v1/portal/hospitals/{hospital_id}/doctors/{doc['id']}/slots/",
               token=tk, query={"date": str(d)})
        bk = r.call("Appointments", who, "Book an appointment online", "post", "/api/v1/portal/book/", token=tk,
                    body={"hospital_id": hospital_id, "doctor_id": doc["id"], "scheduled_date": str(d), "scheduled_time": f"10:{15 * (len(self._booked)):02d}",
                          "chief_complaint": "Fever and cough for two days", "data_sharing_consent": True})
        self._booked.append(1)
        pt["booking"] = bk
        r.call("Appointments", who, "My bookings", "get", "/api/v1/portal/my-bookings/", token=tk)
        fam = r.call("Family", who, "Add a family member", "post", "/api/v1/portal/family/", token=tk,
                     body={"full_name": f"Baby {n[-4:]}", "date_of_birth": str(date.today() - timedelta(days=200)), "gender": "male",
                           "relationship": "child"})
        pt["child"] = (fam or {}).get("awpid")
        r.call("Family", who, "List family members", "get", "/api/v1/portal/family/", token=tk)
        r.call("Patient portal", who, "Health summary", "get", "/api/v1/portal/health-summary/", token=tk)
        r.call("Patient portal", who, "Vaccination roadmap", "get", "/api/v1/portal/vaccinations/", token=tk)
        r.call("Patient portal", who, "Growth", "get", "/api/v1/portal/growth/", token=tk)
        r.call("Patient portal", who, "Milestones", "get", "/api/v1/portal/milestones/", token=tk)
        r.call("Patient portal", who, "Invoices", "get", "/api/v1/portal/invoices/", token=tk)
        r.call("Patient portal", who, "Notifications", "get", "/api/v1/portal/notifications/", token=tk)
        r.call("Patient portal", who, "Emergency token", "post", "/api/v1/portal/emergency/token/", token=tk, body={"consent_confirmed": True})

    def hospital_visit(self, r, pats, staff, adm, branch_id, tag, mob):
        fd, nu, dr = staff["front_desk"][0]["token"], staff["nurse"][0]["token"], staff["doctor"][0]["token"]
        lab, ph = staff["lab_tech"][0]["token"], staff["pharmacist"][0]["token"]
        today = str(date.today())
        pt = pats[0]
        # front desk: find the online booking, check the patient in
        ap = r.call("Front desk", "front_desk", "Upcoming appointments (the online booking is there)", "get", "/api/v1/opd/appointments/upcoming/", token=fd,
                    check=lambda d: True if (d.get("results") if isinstance(d, dict) else d) else "no appointments")
        rows = (ap.get("results") if isinstance(ap, dict) else ap) or []
        appt = next((x for x in rows if x.get("patient_awpid") == pt["awpid"]), rows[0] if rows else {})
        aid = appt.get("id")
        r.call("Front desk", "front_desk", "OPD stats", "get", "/api/v1/opd/stats/", token=fd)
        r.call("Front desk", "front_desk", "Find a patient by mobile", "get", "/api/v1/patients/lookup/", token=fd, query={"mobile": pt["mobile"]})
        r.call("Front desk", "front_desk", "Search patients", "get", "/api/v1/patients/search/", token=fd, query={"q": "Asha"})
        # walk-in
        ty = r.call("Front desk", "front_desk", "Appointment types", "get", "/api/v1/opd/appointment-types/", token=fd)
        tyrows = ty if isinstance(ty, list) else (ty or {}).get("results", [])
        atype = (tyrows[0].get("value") or tyrows[0].get("name")) if tyrows else "consultation"
        wm = mob()
        wi = r.call("Front desk", "front_desk", "Register a walk-in patient", "post", "/api/v1/patients/register/", token=fd,
                    body={"full_name": f"Walk In {tag}", "mobile": wm, "branch_id": branch_id, "dpdp_consent": True, "gender": "M",
                          "date_of_birth": str(date.today() - timedelta(days=3650)), "share_with_network": True, "hie_consent": True})
        wid = (wi or {}).get("uuid")
        if wid:
            r.call("Appointments", "front_desk", "Book a walk-in appointment", "post", "/api/v1/opd/appointments/", token=fd,
                   body={"patient_id": wid, "doctor_id": appt.get("doctor_id"),
                         "scheduled_date": today, "chief_complaint": "Stomach ache", "appointment_type": atype})
        if not aid:
            return
        r.call("Appointments", "front_desk", "Check the patient in (waiting)", "post", f"/api/v1/opd/appointments/{aid}/status/", token=fd, body={"status": "waiting"})
        r.call("Appointments", "front_desk", "Reschedule is rejected for a checked-in visit or accepted", "post", f"/api/v1/opd/appointments/{aid}/reschedule/", token=fd,
               body={"scheduled_date": today}, expect=(200, 400))
        # nurse
        r.call("Nurse", "nurse", "Queue/monitoring view", "get", "/api/v1/opd/monitoring/", token=nu)
        r.call("Nurse", "nurse", "Record vitals (weight, height, temperature)", "post", f"/api/v1/opd/appointments/{aid}/vitals/", token=nu,
               body={"weight_kg": "9.8", "height_cm": "78", "temperature": "100.2", "pulse_rate": 110, "spo2": 98, "head_circumference_cm": "46"})
        # doctor
        r.call("Doctor", "doctor", "Start the consultation", "post", f"/api/v1/opd/appointments/{aid}/status/", token=dr, body={"status": "in_progress"})
        enc = r.call("Doctor", "doctor", "Write an encounter (SOAP, diagnosis)", "post", "/api/v1/opd/encounters/", token=dr,
                     body={"appointment_id": aid, "patient_id": appt.get("patient_id"),
                           "subjective": "Fever and cough", "objective": "Chest clear", "assessment": "Viral URTI", "plan": "Supportive care",
                           "advice_to_patient": "Fluids and rest", "follow_up_in_days": 5,
                           "diagnoses": [{"code": "J06.9", "description": "Acute upper respiratory infection"}]})
        eid = (enc or {}).get("id")
        cat = r.call("Lab", "lab_tech", "Add a test to the lab catalogue", "post", "/api/v1/lab/catalog/", token=lab,
                     body={"name": "CBC", "code": "CBC", "price": "300", "turnaround_hours": 4}, expect=(200, 201))
        rx = r.call("Doctor", "doctor", "Create a prescription", "post", "/api/v1/opd/prescriptions/", token=dr, body={"encounter_id": eid})
        rid = (rx or {}).get("id")
        self.ctx = {"patient_pk": appt.get("patient_pk"), "appt": appt, "rid": rid, "eid": eid}
        if rid:
            r.call("Doctor", "doctor", "Add a medicine to it", "post", f"/api/v1/opd/prescriptions/{rid}/items/", token=dr,
                   body={"drug_name": "Paracetamol", "dosage": "250 mg", "frequency": "td", "duration_days": 3, "quantity": 9, "route": "oral"})
        if cat and cat.get("id"):
            r.call("Doctor", "doctor", "Order a lab test", "post", "/api/v1/lab/requests/", token=dr, body={"encounter_id": eid, "tests": [cat["id"]]})
        r.call("Doctor", "doctor", "Sign the encounter", "post", f"/api/v1/opd/encounters/{eid}/sign/", token=dr, body={})
        r.call("Doctor", "doctor", "Encounter summary PDF", "get", f"/api/v1/opd/encounters/{eid}/pdf/", token=dr)
        r.call("Doctor", "doctor", "Finish the appointment", "post", f"/api/v1/opd/appointments/{aid}/status/", token=dr, body={"status": "done"}, expect=(200, 400))
        # lab
        lr = r.call("Lab", "lab_tech", "Pending lab requests", "get", "/api/v1/lab/requests/", token=lab)
        lrows = (lr.get("results") if isinstance(lr, dict) else lr) or []
        if lrows:
            r.call("Lab", "lab_tech", "Move a lab request to processing", "patch", f"/api/v1/lab/requests/{lrows[0]['id']}/status/", token=lab, body={"status": "processing"})
        # pharmacy
        r.call("Pharmacy", "pharmacist", "Prescriptions waiting", "get", "/api/v1/pharmacy/prescriptions/", token=ph)
        r.call("Pharmacy", "pharmacist", "Stock list", "get", "/api/v1/pharmacy/stock/", token=ph)
        # billing
        inv = r.call("Billing", "front_desk", "Invoices", "get", "/api/v1/billing/invoices/", token=fd)
        r.call("Billing", "hospital admin", "Billing summary for the visit", "get", "/api/v1/billing/summary/", token=adm, query={"appointment_id": aid})
        r.call("Billing", "hospital admin", "Revenue report", "get", "/api/v1/billing/reports/revenue/", token=adm)
        # patient comes back
        tk, who = pt["token"], f"patient {pt['awpid'][-4:]}"
        r.call("Return visit", who, "Visit now in my records", "get", "/api/v1/portal/my-records/", token=tk)
        r.call("Return visit", who, "Prescriptions from the hospital", "get", "/api/v1/portal/prescriptions/", token=tk)
        r.call("Return visit", who, "Lab orders", "get", "/api/v1/portal/lab-orders/", token=tk)
        r.call("Return visit", who, "My bookings after the visit", "get", "/api/v1/portal/my-bookings/", token=tk)
        r.call("Return visit", who, "Timeline after the visit (vitals feed growth)", "get", f"/api/v1/peds/patients/{pt['awpid']}/growth/", token=tk)

    def more_flows(self, r, pats, staff, adm, branch_id, hospital_id, doc, dept_id):
        ctx = getattr(self, "ctx", {})
        pt = pats[0]
        fd, dr, nu, ph = (staff[k][0]["token"] for k in ("front_desk", "doctor", "nurse", "pharmacist"))
        who = f"patient {pt['awpid'][-4:]}"
        pk = ctx.get("patient_pk")
        if not pk:
            found = r.call("Front desk", "front_desk", "Find the patient's record id by AWPID", "get", "/api/v1/patients/search/", token=fd, query={"q": pt["awpid"]})
            rows = (found.get("results") if isinstance(found, dict) else found) or []
            pk = (rows[0] or {}).get("id") if rows else None

        # a parent books for the child
        child = pt.get("child")
        if child:
            day = str(date.today() + timedelta(days=3))
            r.call("Child booking", who, "Parent books a visit for the child", "post", "/api/v1/portal/book/", token=pt["token"],
                   body={"hospital_id": hospital_id, "doctor_id": doc["id"], "scheduled_date": day, "scheduled_time": "12:00",
                         "chief_complaint": "Vaccination due", "patient_awpid": child, "data_sharing_consent": True}, expect=(201,))
            ap = r.call("Child booking", "front_desk", "Front desk sees the child's appointment", "get", "/api/v1/opd/appointments/upcoming/", token=fd)
            rows = (ap.get("results") if isinstance(ap, dict) else ap) or []
            kid = next((x for x in rows if x.get("patient_awpid") == child), None)
            r.note("Child booking", "front_desk", "Child shows as a dependent with the guardian relation", bool(kid) and kid.get("is_dependent") is True
                   and kid.get("guardian_relation") == "child", str({k: (kid or {}).get(k) for k in ("is_dependent", "guardian_name", "guardian_relation")}))
            r.call("Child booking", who, "Parent sees the child's booking", "get", "/api/v1/portal/my-bookings/", token=pt["token"], query={"patient_awpid": child})

        # billing: invoice, partial payment, refused payments, final payment
        if pk:
            inv = r.call("Billing", "front_desk", "Create an invoice with one line", "post", "/api/v1/billing/invoices/", token=fd,
                         body={"patient": pk, "items": [{"description": "Consultation", "quantity": 1, "unit_price": "500"}]}, expect=(200, 201))
            iid = (inv or {}).get("id")
            if iid:
                pay = lambda what, amt, ref="", exp=(201,): r.call("Billing", "front_desk", what, "post", "/api/v1/billing/payments/", token=fd,
                                                                  body={"invoice": iid, "amount": amt, "payment_mode": "cash", "transaction_ref": ref}, expect=exp)
                pay("Part payment of 200", "200", "R1")
                pay("Paying more than is owed is refused", "1000", "R2", (400,))
                pay("The same reference twice on one invoice is refused", "50", "R1", (400,))
                pay("A zero payment is refused", "0", "", (400,))
                pay("Final payment of 300", "300", "R3")
                r.call("Billing", "front_desk", "Invoice is now paid in full", "get", f"/api/v1/billing/invoices/{iid}/", token=fd,
                       check=lambda d: True if d.get("status") == "paid" and str(d.get("paid_amount")) in ("500", "500.00") else f"status={d.get('status')} paid={d.get('paid_amount')}")

        # pharmacy: drug, stock, dispense, over-dispense
        drug = r.call("Pharmacy", "pharmacist", "Add a drug to the catalogue", "post", "/api/v1/prescriptions/drugs/", token=ph,
                      body={"name": "Paracetamol 250", "generic_name": "paracetamol", "form": "Tablet", "strength": "250mg"}, expect=(200, 201))
        did = (drug or {}).get("id")
        if did:
            stock = r.call("Pharmacy", "pharmacist", "Receive 50 units into stock", "post", "/api/v1/pharmacy/stock/", token=ph,
                           body={"drug": did, "branch": branch_id, "batch_number": "B1", "quantity": 50, "mrp": "2", "unit_cost": "1", "expiry_date": "2028-01-01"},
                           expect=(200, 201))
            sid = (stock or {}).get("id")
            rx = r.call("Pharmacy", "doctor", "Read the prescription written earlier", "get", f"/api/v1/opd/prescriptions/{ctx.get('rid')}/", token=dr) if ctx.get("rid") else {}
            item = ((rx or {}).get("items") or [{}])[0].get("id")
            if sid and item:
                r.call("Pharmacy", "pharmacist", "Dispense 3 units against the prescription", "post", "/api/v1/pharmacy/dispense/", token=ph,
                       body={"prescription_item": item, "stock": sid, "quantity": 3}, expect=(200, 201))
                r.call("Pharmacy", "pharmacist", "Dispensing more than is in stock is refused", "post", "/api/v1/pharmacy/dispense/", token=ph,
                       body={"prescription_item": item, "stock": sid, "quantity": 500}, expect=(400,))
                r.call("Pharmacy", "pharmacist", "Stock fell to 47", "get", "/api/v1/pharmacy/stock/", token=ph,
                       check=lambda d: True if any(x.get("quantity") == 47 for x in ((d.get("results") if isinstance(d, dict) else d) or [])) else "no stock row at 47")

        # vaccination actions by staff
        if pk:
            r.call("Vaccination", "doctor", "Doctor orders a vaccine", "post", f"/api/v1/patients/{pk}/vaccinations/order/", token=dr,
                   body={"vaccine_name": "MMR", "reason": "due"}, expect=(200, 201))
            r.call("Vaccination", "nurse", "Nurse marks a dose given", "post", f"/api/v1/patients/{pk}/vaccinations/administer/", token=nu,
                   body={"vaccine_name": "BCG", "dose_number": 1}, expect=(200, 201))
            r.call("Vaccination", "doctor", "Doctor marks a slot not required", "post", f"/api/v1/patients/{pk}/vaccinations/decline/", token=dr,
                   body={"vaccine_name": "Hepatitis B", "reason": "given elsewhere"}, expect=(200, 201))
            r.call("Vaccination", "doctor", "Vaccination schedule view", "get", f"/api/v1/patients/{pk}/vaccinations/", token=dr)

        # inpatient: referral, registration, bed, discharge
        types = r.call("Admission", "front_desk", "Admission types configured", "get", "/api/v1/ipd/admission-types/", token=fd)
        srcs = r.call("Admission", "front_desk", "Admission sources configured", "get", "/api/v1/ipd/admission-sources/", token=fd)
        first = lambda d: ((d if isinstance(d, list) else (d or {}).get("results")) or [{}])[0].get("value")
        if pk and dept_id and first(types) and first(srcs):
            rf = r.call("Admission", "doctor", "Doctor recommends admission", "post", "/api/v1/ipd/referrals/recommend/", token=dr,
                        body={"patient_id": pk, "department_id": dept_id, "admission_type": first(types), "admission_source": first(srcs),
                              "reason_for_admission": "Needs observation"}, expect=(200, 201))
            rid2 = (rf or {}).get("id")
            rtypes = r.call("Admission", "hospital admin", "Room types", "get", "/api/v1/billing/room-types/", token=adm)
            bedtype = next((x.get("value") for x in ((rtypes if isinstance(rtypes, list) else (rtypes or {}).get("results")) or []) if x.get("is_bed_based")), None)
            fls = r.call("Admission", "hospital admin", "Floors at the branch", "get", "/api/v1/org/floors/", token=adm, query={"branch_id": branch_id})
            flrows = fls if isinstance(fls, list) else (fls or {}).get("results", [])
            floor_id = flrows[0]["id"] if flrows else (r.call(
                "Admission", "hospital admin", "Create a floor", "post", "/api/v1/org/floors/", token=adm,
                body={"branch": branch_id, "name": "Ground"}, expect=(200, 201)) or {}).get("id")
            room = r.call("Admission", "hospital admin", "Create a ward room", "post", "/api/v1/org/rooms/", token=adm,
                          body={"branch": branch_id, "floor_obj": floor_id, "name": "Ward 1", "room_type": bedtype or "general_ward", "capacity": 4}, expect=(200, 201))
            bed = r.call("Admission", "hospital admin", "Create a bed", "post", "/api/v1/org/beds/", token=adm,
                         body={"room": (room or {}).get("id"), "bed_number": "1"}, expect=(200, 201))
            if rid2:
                ad = r.call("Admission", "front_desk", "Front desk registers the admission", "post", "/api/v1/ipd/admissions/register/", token=fd,
                            body={"referral_id": rid2, "attendant_name": "Parent", "attendant_phone": "9000000002", "consent_given": True}, expect=(200, 201))
                aid = (ad or {}).get("id")
                if aid and (bed or {}).get("id"):
                    r.call("Admission", "front_desk", "Assign the bed", "post", f"/api/v1/ipd/admissions/{aid}/assign-bed/", token=fd, body={"bed_id": bed["id"]}, expect=(200, 201))
                    r.call("Admission", "front_desk", "Discharge", "post", f"/api/v1/ipd/admissions/{aid}/discharge/", token=fd, body={}, expect=(200, 201))

    def access_control(self, r, staff, adm, plat, pats, logins, hospital_id):
        nu, fd, ph, lab = (staff[k][0]["token"] for k in ("nurse", "front_desk", "pharmacist", "lab_tech"))
        pt = pats[0]["token"]
        r.call("Access control", "nurse", "Nurse cannot invite staff", "post", "/api/v1/org/staff/invite/", token=nu,
               body={"first_name": "X", "phone": "9000000001", "role": "nurse"}, expect=(401, 403))
        r.call("Access control", "pharmacist", "Pharmacist cannot create a lab test", "post", "/api/v1/lab/catalog/", token=ph,
               body={"name": "X"}, expect=(401, 403))
        r.call("Access control", "front_desk", "Front desk cannot write an encounter", "post", "/api/v1/opd/encounters/", token=fd, body={}, expect=(401, 403))
        r.call("Access control", "patient", "Patient cannot open the staff list", "get", "/api/v1/org/staff/", token=pt, expect=(401, 403))
        r.call("Access control", "patient", "Patient cannot open the platform", "get", "/api/v1/platform/hospitals/", token=pt, expect=(401, 403))
        r.call("Access control", "hospital admin", "Hospital admin cannot open the platform", "get", "/api/v1/platform/hospitals/", token=adm, expect=(401, 403))
        r.call("Access control", "anonymous", "No token is refused", "get", "/api/v1/portal/profile/", expect=(401, 403))
        other = pats[1]["awpid"] if len(pats) > 1 else None
        if other:
            r.call("Access control", "patient", "A patient cannot read another patient's timeline", "get", f"/api/v1/peds/patients/{other}/health-timeline/",
                   token=pt, expect=(401, 403, 404))
        # platform-level document tools
        r.call("Documents", "platform admin", "Document report across patients", "get", "/api/v1/platform/records/report/", token=plat)
        r.call("Documents", "platform admin", "Re-run the classification rules", "post", "/api/v1/platform/records/reclassify/", token=plat, body={})
        # auth housekeeping
        lg = r.call("Login", "patient", "Patient OTP login: request", "post", "/api/v1/auth/otp/request/", body={"purpose": "login_patient", "identifier": pats[0]["awpid"]},
                    expect=(200, 201, 400))
        r.call("Login", "patient", "Token refresh", "post", "/api/v1/auth/token/refresh/", body={"refresh": "bad"}, expect=(400, 401))
        r.call("Login", "doctor", "Wrong password is refused", "post", "/api/v1/auth/login/staff/", body={"mobile": staff["doctor"][0]["mobile"], "password": "nope-nope"}, expect=(400, 401))
        r.call("Login", "hospital admin", "Staff list after all invites", "get", "/api/v1/org/staff/", token=adm,
               check=lambda d: True if len((d.get("results") if isinstance(d, dict) else d) or []) >= 6 else f"only {len((d.get('results') if isinstance(d, dict) else d) or [])} staff: {str(d)[:80]}")
        with open(f"logs/e2e_logins_{r.tag}.txt", "w", encoding="utf-8") as fh:
            fh.write(f"hospital_code: e2e-{r.tag}\n")
            for role, ident, pw in logins:
                fh.write(f"{role}\t{ident}\t{pw}\n")

    def documents(self, r, pt):
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from django.core.files.uploadedfile import SimpleUploadedFile
        pages = {
            "prescription.pdf": ["Prescription", "Rx", "Tab Paracetamol 250 mg twice daily for 3 days", "Syrup Amoxicillin 5 ml thrice daily",
                                 "Diagnosis: acute upper respiratory infection", "Advice: plenty of fluids. Review after 5 days.", "Dr. Meera Nair, Pediatrics"],
            "lab_report.pdf": ["City Diagnostics Laboratory - Laboratory Report", "Complete Blood Count", "Hemoglobin 10.2 g/dL  Reference range 11.5-14.5",
                               "WBC count 11200 /cumm  Reference range 6000-14000", "Platelet count 310000  Reference range 150000-450000",
                               "Fasting glucose 82 mg/dL  Reference range 70-100", "Specimen: blood. Result reviewed by pathologist."],
            "discharge.pdf": ["Discharge Summary", "Date of admission: 02-03-2026   Date of discharge: 06-03-2026", "Ward: Pediatric ward",
                              "Diagnosis: bronchiolitis with hypoxia", "Treatment given: oxygen, nebulisation, IV fluids", "Condition at discharge: stable",
                              "Follow up after one week."],
            "growth.pdf": ["Growth Chart - Well Baby Clinic", "Date of visit: 10-03-2026", "Weight 7.4 kg", "Height 66 cm", "Head circumference 42.5 cm",
                           "Age 7 months. Development appropriate for age."],
        }
        files = []
        for name, lines in pages.items():
            buf = io.BytesIO()
            c = canvas.Canvas(buf, pagesize=A4)
            y = 800
            for line in lines:
                c.drawString(60, y, line)
                y -= 24
            c.save()
            files.append(SimpleUploadedFile(name, buf.getvalue(), content_type="application/pdf"))
        tk, who = pt["token"], f"patient {pt['awpid'][-4:]}"
        up = r.call("Documents", who, "Upload four documents at once (prescription, lab, discharge, growth)", "post", "/api/v1/records/upload/",
                    token=tk, body={}, files={"files": files}, expect=(200, 201, 202))
        print("UPLOAD", json.dumps(up, default=str)[:500])
        r.call("Documents", who, "Document list with classification", "get", "/api/v1/portal/documents/", token=tk)
        rows = up.get("documents", []) if isinstance(up, dict) else []
        need = [x for x in rows if x.get("processing_status") == "review_required"]
        if need:
            r.call("Documents", who, "Review step: patient picks the type of a document the rules could not place", "post", "/api/v1/records/submit/",
                   token=tk, body={"decisions": [{"document_id": x["id"], "document_type": "growth_chart"} for x in need]}, expect=(200, 201, 202, 400))
        r.call("Documents", who, "Document types offered", "get", "/api/v1/records/types/", token=tk)
        if up.get("batch_id"):
            r.call("Documents", who, "Batch status", "get", f"/api/v1/records/batches/{up['batch_id']}/", token=tk)
        r.call("Documents", who, "Document counts by type", "get", "/api/v1/portal/documents/counts/", token=tk)
        tl = r.call("Documents", who, "Insights timeline built from the documents", "get", f"/api/v1/peds/patients/{pt['awpid']}/health-timeline/", token=tk)
        print("TIMELINE", json.dumps(tl, default=str)[:400])

    def excel(self, r, path):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Tested"
        head = ["Area", "Login", "Functionality tested", "Request", "Expected", "Got", "Result", "Note"]
        ws.append(head)
        for row in r.rows:
            ws.append(row)
        arial = Font(name="Arial", size=10)
        for c in ws[1]:
            c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = arial
                c.alignment = Alignment(vertical="top", wrap_text=c.column in (3, 4, 8))
            res = row[6]
            res.fill = PatternFill("solid", fgColor="C6EFCE" if res.value == "PASS" else "FFC7CE")
        for col, w in zip("ABCDEFGH", (16, 16, 48, 52, 10, 8, 9, 50)):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A2"
        n = len(r.rows) + 1
        s = wb.create_sheet("Summary", 0)
        s.append(["Measure", "Value"])
        s.append(["Checks run", f"=COUNTA(Tested!G2:G{n})"])
        s.append(["Passed", f'=COUNTIF(Tested!G2:G{n},"PASS")'])
        s.append(["Failed", f'=COUNTIF(Tested!G2:G{n},"FAIL")'])
        s.append(["Pass rate", "=IF(B2=0,0,B3/B2)"])
        s["B5"].number_format = "0.0%"
        s.append([])
        s.append(["Area", "Checks", "Passed"])
        for i, area in enumerate(dict.fromkeys(x[0] for x in r.rows), start=8):
            s.append([area, f'=COUNTIF(Tested!A2:A{n},A{i})', f'=COUNTIFS(Tested!A2:A{n},A{i},Tested!G2:G{n},"PASS")'])
        for row in s.iter_rows():
            for c in row:
                c.font = arial
        for c in (s["A1"], s["B1"], s["A7"], s["B7"], s["C7"]):
            c.font = Font(name="Arial", size=10, bold=True)
        s.column_dimensions["A"].width = 24
        wb.save(path)
